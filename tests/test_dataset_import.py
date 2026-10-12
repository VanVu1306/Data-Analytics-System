"""Test importer CSV / Excel / SQLite và DatasetStore (không cần LLM, không cần mạng)."""

import datetime
import io
import sqlite3
import zipfile

import pytest
from openpyxl import Workbook

from backend.database.dataset_store import DatasetStore
from backend.database.importers import ImportLimits
from backend.errors import (
    DatasetError,
    DatasetNotFoundError,
    FileTooLargeError,
    InvalidFileError,
    UnsafeContentError,
    UnsupportedFormatError,
)


@pytest.fixture
def store(tmp_path):
    return DatasetStore(tmp_path / "datasets")


def _rows(store, info, table):
    conn = sqlite3.connect(store.db_path(info.dataset_id))
    try:
        return conn.execute(f'SELECT * FROM "{table}"').fetchall()
    finally:
        conn.close()


def _xlsx(build) -> bytes:
    wb = Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _sqlite_bytes(tmp_path, script: str) -> bytes:
    path = tmp_path / f"src_{len(list(tmp_path.glob('src_*.db')))}.db"
    conn = sqlite3.connect(path)
    conn.executescript(script)
    conn.commit()
    conn.close()
    return path.read_bytes()


def _leftovers(store):
    return list(store.root.iterdir())


# ----------------------------- CSV ------------------------------------------ #
def test_csv_types_and_normalized_names(store):
    csv = "Mã SP;Tên;Giá;SL;Ghi chú\n007;Áo;12.5;3;\n010;Quần;9;4;ok\n".encode("utf-8-sig")
    info = store.create("Doanh thu 2026.csv", csv)
    (table,) = info.tables
    assert table["name"] == "Doanh_thu_2026" and table["source_name"] == "Doanh thu 2026"
    types = {c["name"]: c["type"] for c in table["columns"]}
    assert types == {"Mã_SP": "TEXT", "Tên": "TEXT", "Giá": "REAL", "SL": "INTEGER", "Ghi_chú": "TEXT"}
    rows = _rows(store, info, "Doanh_thu_2026")
    assert rows[0] == ("007", "Áo", 12.5, 3, None)  # số 0 đứng đầu giữ nguyên, ô rỗng -> NULL
    assert table["row_count"] == 2


def test_csv_duplicate_and_empty_headers(store):
    info = store.create("d.csv", b"a,a,,1x\n1,2,3,4\n")
    assert [c["name"] for c in info.tables[0]["columns"]] == ["a", "a_2", "column_3", "c_1x"]


def test_csv_header_only_gives_empty_table(store):
    info = store.create("h.csv", b"a,b\n")
    assert info.tables[0]["row_count"] == 0


def test_csv_windows_1258_warns(store):
    info = store.create("v.csv", "tên,giá\nÁo,1\n".encode("cp1258"))
    assert any("Windows-1258" in w for w in info.warnings)
    assert _rows(store, info, "v")[0][0] == "Áo"


def test_csv_rejects_binary_empty_and_too_many_rows(tmp_path):
    store = DatasetStore(tmp_path / "d", ImportLimits(max_rows_per_table=3))
    with pytest.raises(InvalidFileError):
        store.create("a.csv", b"PK\x03\x04binary")
    with pytest.raises(InvalidFileError):
        store.create("a.csv", b"")
    with pytest.raises(InvalidFileError):
        store.create("a.csv", b"\n\n  \n")
    with pytest.raises(InvalidFileError):
        store.create("a.csv", b"a\n1\n2\n3\n4\n")
    assert _leftovers(store) == []


# ----------------------------- Excel ---------------------------------------- #
def _multi_sheet(wb):
    ws = wb.active
    ws.title = "Đơn hàng"
    ws.append(["BÁO CÁO THÁNG 1"])
    ws.append(["id", "ngày", "tiền", "ghi chú"])
    ws.append([1, datetime.datetime(2026, 1, 2), 10.5, "ok"])
    ws.append([2, datetime.datetime(2026, 1, 3), 3, "#N/A"])
    kh = wb.create_sheet("Khách hàng")
    kh.append(["id", "tên", None, "id"])
    kh.append([1, "An", None, "x"])
    wb.create_sheet("Trống")
    hidden = wb.create_sheet("Ẩn")
    hidden.append(["a"])
    hidden.append([1])
    hidden.sheet_state = "hidden"
    wb.create_sheet("Chỉ tiêu đề").append(["a", "b"])


def test_excel_sheets_become_tables_in_one_database(store):
    info = store.create("data.xlsx", _xlsx(_multi_sheet))
    assert [t["name"] for t in info.tables] == ["Đơn_hàng", "Khách_hàng"]
    don = info.tables[0]
    assert {c["name"]: c["type"] for c in don["columns"]} == {
        "id": "INTEGER",
        "ngày": "TEXT",
        "tiền": "REAL",
        "ghi_chú": "TEXT",
    }
    assert _rows(store, info, "Đơn_hàng") == [(1, "2026-01-02", 10.5, "ok"), (2, "2026-01-03", 3.0, None)]
    assert [c["name"] for c in info.tables[1]["columns"]] == ["id", "tên", "id_2"]  # cột rỗng bị bỏ, tên trùng được đánh số
    joined = " | ".join(info.warnings)
    for expected in ("bỏ qua 1 dòng phía trên", "'Trống'", "'Ẩn'", "'Chỉ tiêu đề'"):
        assert expected in joined


def test_excel_mixed_column_becomes_text_and_int_float_becomes_real(store):
    def build(wb):
        ws = wb.active
        ws.title = "S"
        ws.append(["a", "b"])
        ws.append([1, 1.5])
        ws.append(["x", 5])

    info = store.create("m.xlsx", _xlsx(build))
    assert {c["name"]: c["type"] for c in info.tables[0]["columns"]} == {"a": "TEXT", "b": "REAL"}
    assert _rows(store, info, "S") == [("1", 1.5), ("x", 5.0)]


def test_infer_column_edge_cases():
    from backend.database.importers import infer_column

    assert infer_column([1, None, 10**20]) == ("TEXT", ["1", None, "100000000000000000000"])  # vượt int64
    assert infer_column([True, False, None]) == ("INTEGER", [1, 0, None])
    assert infer_column([1.5, float("nan"), 2]) == ("REAL", [1.5, None, 2.0])
    assert infer_column([datetime.datetime(2026, 1, 2, 3, 4)]) == ("TEXT", ["2026-01-02 03:04:00"])
    assert infer_column([None, None]) == ("TEXT", [None, None])


def test_excel_rejects_bad_files(tmp_path):
    store = DatasetStore(tmp_path / "d", ImportLimits(max_rows_per_table=5))
    with pytest.raises(InvalidFileError):  # không phải zip
        store.create("a.xlsx", b"not an excel file")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:  # zip nhưng không phải workbook
        zf.writestr("hello.txt", "x")
    with pytest.raises(InvalidFileError):
        store.create("a.xlsx", buf.getvalue())

    def many(wb):
        ws = wb.active
        ws.append(["a"])
        for i in range(50):
            ws.append([i])

    with pytest.raises(InvalidFileError):  # quá số dòng
        store.create("a.xlsx", _xlsx(many))
    with pytest.raises(InvalidFileError):  # workbook không có dữ liệu
        store.create("a.xlsx", _xlsx(lambda wb: None))
    assert _leftovers(store) == []


def test_excel_zip_bomb_is_rejected(tmp_path):
    store = DatasetStore(tmp_path / "d", ImportLimits(max_uncompressed_bytes=1_000_000))
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("xl/workbook.xml", "<x/>")
        zf.writestr("xl/big.bin", b"0" * 5_000_000)  # nén rất nhỏ, giải nén vượt giới hạn
    assert len(bomb.getvalue()) < 100_000
    with pytest.raises(FileTooLargeError):
        store.create("a.xlsx", bomb.getvalue())
    assert _leftovers(store) == []


# ----------------------------- SQLite --------------------------------------- #
SCRIPT = """
CREATE TABLE a(id INTEGER PRIMARY KEY, n TEXT);
CREATE TABLE b(id INTEGER PRIMARY KEY, a_id INT REFERENCES a(id), v REAL,
               g INT GENERATED ALWAYS AS (id * 2) VIRTUAL);
CREATE VIEW vv AS SELECT * FROM a;
CREATE TRIGGER tg AFTER INSERT ON a BEGIN DELETE FROM b; END;
CREATE INDEX ix ON b(v);
INSERT INTO a VALUES (1, 'x'), (2, 'y');
INSERT INTO b(id, a_id, v) VALUES (1, 1, 2.5);
"""


def test_sqlite_upload_is_rebuilt_not_copied(store, tmp_path):
    info = store.create("my data.sqlite", _sqlite_bytes(tmp_path, SCRIPT))
    assert [(t["name"], t["row_count"]) for t in info.tables] == [("a", 2), ("b", 1)]
    assert info.relations == [{"from_table": "b", "from_column": "a_id", "to_table": "a", "to_column": "id"}]
    assert any("view" in w for w in info.warnings) and any("trigger" in w for w in info.warnings)
    conn = sqlite3.connect(store.db_path(info.dataset_id))
    try:
        kinds = {r[0] for r in conn.execute("SELECT type FROM sqlite_master")}
        assert kinds == {"table"}  # không có trigger/view/index nào chạy mã về sau
        assert conn.execute("SELECT * FROM b").fetchall() == [(1, 1, 2.5, 2)]
    finally:
        conn.close()


def test_sqlite_rejects_bad_files(store, tmp_path):
    with pytest.raises(InvalidFileError):
        store.create("a.db", b"this is not sqlite at all")
    good = _sqlite_bytes(tmp_path, SCRIPT)
    with pytest.raises(DatasetError):  # hỏng ở giữa file
        store.create("a.db", good[:100] + b"\xff" * 3000)
    with pytest.raises(InvalidFileError):  # hợp lệ nhưng không có bảng
        store.create("a.db", _sqlite_bytes(tmp_path, "CREATE VIEW v AS SELECT 1;"))
    assert _leftovers(store) == []


def test_sqlite_dangerous_function_in_schema_is_blocked(store, tmp_path):
    script = "CREATE TABLE t(a INT CHECK (load_extension('x') IS NULL OR 1)); INSERT INTO t VALUES (1);"
    with pytest.raises(UnsafeContentError):
        store.create("a.db", _sqlite_bytes(tmp_path, script))
    assert _leftovers(store) == []
    harmless = "CREATE TABLE t(a INT CHECK (abs(a) >= 0)); INSERT INTO t VALUES (1);"
    assert store.create("ok.db", _sqlite_bytes(tmp_path, harmless)).tables[0]["row_count"] == 1


# ----------------------------- Store ---------------------------------------- #
@pytest.mark.parametrize("name", ["a.sql", "a.json", "a.xls", "a.txt", "noext"])
def test_unsupported_formats(store, name):
    with pytest.raises(UnsupportedFormatError):
        store.create(name, b"x")


def test_sql_script_message_explains_why(store):
    with pytest.raises(UnsupportedFormatError, match="script"):
        store.create("dump.sql", b"DROP TABLE x;")


def test_size_limit_and_db_limit(tmp_path):
    store = DatasetStore(tmp_path / "d", ImportLimits(max_upload_bytes=50))
    with pytest.raises(FileTooLargeError):
        store.create("a.csv", b"a\n" + b"1\n" * 100)
    store = DatasetStore(tmp_path / "e", ImportLimits(max_db_bytes=8192))
    big = ("a,b\n" + "".join(f"{i},{'x' * 50}\n" for i in range(2000))).encode()
    with pytest.raises(FileTooLargeError):
        store.create("a.csv", big)
    assert _leftovers(store) == []


def test_each_upload_is_isolated_and_temp_file_removed(store):
    one = store.create("a.csv", b"x\n1\n")
    two = store.create("a.csv", b"x\n2\n3\n")
    assert one.dataset_id != two.dataset_id
    assert store.db_path(one.dataset_id) != store.db_path(two.dataset_id)
    assert _rows(store, one, "a") == [(1,)] and _rows(store, two, "a") == [(2,), (3,)]
    assert sorted(p.name for p in (store.root / one.dataset_id).iterdir()) == ["data.db", "meta.json"]
    assert store.get(one.dataset_id).to_dict() == one.to_dict()


def test_client_filename_never_becomes_a_path(store):
    info = store.create("../../etc/evil.csv", b"x\n1\n")
    assert info.filename == "evil.csv"
    assert [p.name for p in store.root.iterdir()] == [info.dataset_id]


@pytest.mark.parametrize("bad_id", ["../x", "..", "A" * 32, "abc", "", "0" * 31 + "/", None])
def test_untrusted_dataset_id_is_rejected(store, bad_id):
    with pytest.raises(DatasetNotFoundError):
        store.get(bad_id)
    with pytest.raises(DatasetNotFoundError):
        store.db_path(bad_id)


def test_unknown_but_well_formed_id(store):
    with pytest.raises(DatasetNotFoundError):
        store.get("0" * 32)


def test_excel_with_wrong_declared_dimension_is_fully_read(store):
    import re

    wb = Workbook()
    wb.active.append(["a", "b"])
    wb.active.append([1, 2])
    wb.active.append([3, 4])
    raw = io.BytesIO()
    wb.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', data)
            zout.writestr(item, data)
    info = store.create("dim.xlsx", out.getvalue())
    assert info.tables[0]["row_count"] == 2


def test_sqlite_row_limit_enforced(tmp_path):
    path = tmp_path / "big.db"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE t (a)")
    conn.executemany("INSERT INTO t VALUES (?)", [(i,) for i in range(20)])
    conn.commit()
    conn.close()
    small = DatasetStore(tmp_path / "ds", ImportLimits(max_rows_per_table=10))
    with pytest.raises(InvalidFileError):
        small.create("big.db", path.read_bytes())
    assert list(small.root.iterdir()) == []
